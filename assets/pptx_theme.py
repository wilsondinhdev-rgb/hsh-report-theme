"""HSH Report Theme for python-pptx (16:9, 13.333x7.5 in = 960x540 pt).

Usage:
    python3 pptx_theme.py demo out.pptx      # sample deck with the core components
    from pptx_theme import *                 # reuse helpers in your own builder
"""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

def rgb(h): return RGBColor.from_string(h.lstrip('#'))

C = dict(
    red='FF0000', red_deep='D00018', red_bar='CB282D', red_chart='E62601', th_red='CB001C',
    navy='002060', navy_chart='001D5C', navy_banner='092F54', blue='0871D0',
    tint_red='FDEDEC', tint_rose='FBF1F0', tint_icon='FCD8DA', tint_blue='E7F1FB', tint_sky='CCECFF',
    card='F6F7F8', stat='F5F7F9', beige='EEECE1', line='BFC9D9', divider='D9D9D9', row_band='F6F3F3',
)
PILLARS = ['D00018', '0871D0', 'D85A00', '008E52', '87329D', '0098A5', 'D7A000', '334D64']
CHANNELS = ['0076BD', 'E67400', 'C48A00', '009A51', 'D00913']
BSC = dict(fin='FF0000', ops='0070C0', cus='E46C0A', dev='00B050')
FONT = 'Arial'


def new_deck():
    p = Presentation()
    p.slide_width, p.slide_height = Pt(960), Pt(540)
    return p


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def text(slide, x, y, w, h, runs, size=11, bold=False, color='navy', align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, upper=False):
    """runs: str or list of (text, color_key_or_hex[, bold, size])."""
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    lines = runs if isinstance(runs, list) and runs and isinstance(runs[0], list) else [runs]
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        for r in ([(line, color)] if isinstance(line, str) else line):
            t, col = r[0], r[1]; b = r[2] if len(r) > 2 else bold; s = r[3] if len(r) > 3 else size
            run = para.add_run(); run.text = t.upper() if upper else t
            f = run.font; f.name = FONT; f.size = Pt(s); f.bold = b; f.color.rgb = rgb(C.get(col, col))
    return tb


def rect(slide, x, y, w, h, fill, radius=None, line=None, shadow=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    if radius: shp.adjustments[0] = min(0.5, radius / min(w, h))
    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(C.get(fill, fill))
    if line: shp.line.color.rgb = rgb(C.get(line, line)); shp.line.width = Pt(0.75)
    else: shp.line.fill.background()
    if not shadow: shp.shadow.inherit = False
    return shp


def header(slide, crumb1, crumb2, num='03', brand='BRAND NAME'):
    """A1 header strip."""
    text(slide, 36, 27, 45, 30, num, 28, True, 'red')
    rect(slide, 85, 26, 0.75, 28, 'divider')
    text(slide, 102, 27, 400, 12, crumb1, 9, True, 'navy', upper=True)
    text(slide, 102, 40, 400, 12, crumb2, 8, False, 'navy', upper=True)
    text(slide, 640, 30, 260, 16, brand, 12, True, 'red', align=PP_ALIGN.RIGHT)


def title(slide, red_part, navy_part='', y=90, size=26, underline=True):
    """A2 two-tone uppercase title + A3 red accent."""
    runs = [(red_part.upper() + (' ' if navy_part else ''), 'red', True, size)]
    if navy_part: runs.append((navy_part.upper(), 'navy', True, size))
    text(slide, 22, y, 916, 40, [runs])
    if underline: rect(slide, 23, y + 41, 53, 2, 'red')


def badge(slide, x, y, n, w=40, h=32, fill='red_deep', size=14):
    b = rect(slide, x, y, w, h, fill, radius=6)
    tf = b.text_frame; tf.text = n
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    f = p.runs[0].font; f.name = FONT; f.size = Pt(size); f.bold = True; f.color.rgb = rgb('FFFFFF')
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return b


def kpi_tile(slide, x, y, label, value, foot_label, foot_value, foot_unit, blue=False, w=280, h=194):
    """C2 KPI tile."""
    col = 'blue' if blue else 'red'
    rect(slide, x, y, w, h, 'tint_blue' if blue else 'tint_red', radius=14)
    text(slide, x + 15, y + 20, w - 30, 16, label, 12, True, col, upper=True)
    text(slide, x + 15, y + 44, w - 30, 64, value, 54, True, col)
    rect(slide, x + 15, y + h - 50, w - 30, 0.5, 'line')
    text(slide, x + 15, y + h - 38, 100, 24, foot_label, 11, False, 'navy', anchor=MSO_ANCHOR.BOTTOM)
    text(slide, x + 100, y + h - 44, w - 115, 30, [[(foot_value + ' ', 'navy', True, 28), (foot_unit, 'navy', False, 11)]], align=PP_ALIGN.RIGHT)


def stat_card(slide, x, y, value, label, caption, w=275, h=87):
    """C3 stat card with red left bar."""
    rect(slide, x, y, w, h, 'stat', radius=10)
    rect(slide, x, y, 3.2, h, 'red')
    text(slide, x + 22, y + 14, 130, 34, value, 28, True, 'red')
    text(slide, x + 140, y + 20, w - 150, 30, label, 9, True, 'navy', upper=True, align=PP_ALIGN.RIGHT)
    text(slide, x + 22, y + h - 28, w - 40, 16, caption, 9, False, 'navy')


def agenda(slide, labels, shown=None, title_text='GIẢI PHÁP THỰC HIỆN'):
    """B1 agenda grid; shown = number of pillars revealed (progress)."""
    text(slide, 0, 88, 960, 40, title_text, 30, True, 'red', align=PP_ALIGN.CENTER)
    shown = shown or len(labels)
    for i, lab in enumerate(labels[:shown]):
        col = PILLARS[i]; cx = 17 + (i % 4) * 236; cy = 147 + (i // 4) * 177
        rect(slide, cx, cy, 224, 163, 'card', radius=14)
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(cx + 16), Pt(cy + 26), Pt(73), Pt(73))
        dot.fill.solid(); dot.fill.fore_color.rgb = rgb(col); dot.line.fill.background()
        rect(slide, cx + 100, cy + 26, 0.75, 70, col)
        text(slide, cx + 108, cy + 26, 110, 30, '%02d' % (i + 1), 28, True, col)
        text(slide, cx + 108, cy + 66, 110, 50, lab, 12, True, 'navy', upper=True)
        rect(slide, cx + 81, cy + 143, 62, 2.4, col)
    rect(slide, 0, 527.3, 960, 12.7, 'red_deep')


def table(slide, x, y, w, rows, col_w, header_fill='th_red', total_row=False, row_h=20.5, head_h=28.6):
    """F1/F2 table: rows[0] = header."""
    n, m = len(rows), len(rows[0])
    gt = slide.shapes.add_table(n, m, Pt(x), Pt(y), Pt(w), Pt(head_h + row_h * (n - 1))).table
    for j, cw in enumerate(col_w): gt.columns[j].width = Pt(cw)
    for i, r in enumerate(rows):
        gt.rows[i].height = Pt(head_h if i == 0 else row_h)
        for j, v in enumerate(r):
            cell = gt.cell(i, j); cell.margin_left = cell.margin_right = Pt(6); cell.margin_top = cell.margin_bottom = Pt(1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]; run = p.add_run(); run.text = str(v).upper() if i == 0 else str(v)
            f = run.font; f.name = FONT
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = rgb(C.get(header_fill, header_fill)); f.size = Pt(10); f.bold = True; f.color.rgb = rgb('FFFFFF')
                p.alignment = PP_ALIGN.CENTER
            else:
                last = total_row and i == n - 1
                cell.fill.fore_color.rgb = rgb(C['tint_sky'] if last else 'FFFFFF')
                f.size = Pt(9.5); f.bold = last or i == 1 or j == 0
                f.color.rgb = rgb(C['th_red'] if j == 0 else C['navy'])
                p.alignment = PP_ALIGN.CENTER if j == 0 else (PP_ALIGN.RIGHT if str(v).replace('.', '').replace(',', '').isdigit() else PP_ALIGN.LEFT)
    return gt


def combo_chart(slide, x, y, w, h, cats, bars, bar_name='Doanh số'):
    """G1-style bars (navy). Add the red line series in Excel/PowerPoint if a dual axis is needed."""
    cd = CategoryChartData(); cd.categories = cats; cd.add_series(bar_name, bars)
    ch = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(x), Pt(y), Pt(w), Pt(h), cd).chart
    s = ch.series[0]; s.format.fill.solid(); s.format.fill.fore_color.rgb = rgb(C['navy_chart'])
    ch.has_title = False; ch.plots[0].gap_width = 60; ch.has_legend = True; ch.legend.position = XL_LEGEND_POSITION.TOP; ch.legend.include_in_layout = False
    for ax in (ch.category_axis, ch.value_axis):
        ax.tick_labels.font.size = Pt(9); ax.tick_labels.font.name = FONT; ax.tick_labels.font.color.rgb = rgb(C['navy'])
    ch.value_axis.has_major_gridlines = False
    return ch


def demo(path):
    prs = new_deck()
    s = blank(prs); header(s, 'Kế hoạch hoạt động NĐTC 2026 – 2027', 'Giải pháp thực hiện')
    agenda(s, ['Giải pháp kinh doanh', 'Phát triển cửa hàng', 'Tối ưu hệ thống cung ứng', 'Phát triển đa dạng sản phẩm',
               'Hoạt động marketing', 'Đầu tư công nghệ', 'Phát triển nguồn nhân lực', 'Nâng cao hoạt động kiểm soát'])
    s = blank(prs); header(s, 'Kế hoạch hoạt động NĐTC 2026 – 2027', '4. Phát triển đa dạng sản phẩm'); title(s, 'Kế hoạch & lũy kế đến 30/9/2027', y=68)
    data = [('SP thương hiệu công ty', '36', '93', 'sản phẩm', False), ('SP thương hiệu nhà cung cấp', '24', '75', 'sản phẩm', True), ('Tổng sản phẩm mới', '60', '168', 'sản phẩm', False),
            ('Tổng SKU', '2,000', '21,454', 'SKU', False), ('Quốc gia OEM', '04', '10', 'quốc gia', True), ('Hội chợ, triển lãm', '25', '06', 'quốc gia', False)]
    for i, (k, v, fv, fu, b) in enumerate(data):
        kpi_tile(s, 43.6 + (i % 3) * 296, 116.6 + (i // 3) * 206, k, v, 'Lũy kế', fv, fu, blue=b)
    s = blank(prs); header(s, 'Giải pháp & định hướng hoạt động', '1. Giải pháp kinh doanh'); title(s, 'Tình hình khai thác các', 'công ty xây dựng trọn gói')
    for i, (v, k, c) in enumerate([('1.203', 'Đã tiếp xúc', 'Công ty xây dựng trọn gói'), ('230', 'Đã phát sinh giao dịch', 'Công ty đã mua hàng tại HSH'), ('19,1%', 'Tỷ lệ phát sinh doanh số', '230 / 1.203 khách hàng')]):
        stat_card(s, 28.6 + i * 310, 156, v, k, c)
    table(s, 32.5, 293.6, 898, [['STT', 'Tên khách hàng', 'Doanh số (tỷ đồng)'], ['01', 'Công ty A', '1,701'], ['02', 'Công ty B', '173'], ['03', 'Công ty C', '157'], ['', 'Tổng', '2,031']], [55, 680, 163], total_row=True)
    s = blank(prs); header(s, 'Giải pháp & định hướng hoạt động', '1. Giải pháp kinh doanh'); title(s, 'Phát triển', 'đại lý tiêu chuẩn')
    combo_chart(s, 260, 130, 680, 380, ['11/2025', '12/2025', '01/2026', '02/2026', '03/2026', '04/2026'], [8.9, 468.6, 284.1, 382.5, 1451.4, 4264.5])
    prs.save(path); print('saved', path)


if __name__ == '__main__':
    if len(sys.argv) >= 3 and sys.argv[1] == 'demo': demo(sys.argv[2])
    else: print(__doc__)
